# rag/loader.py
"""
Document ingestion: extract text -> split into chunks -> embed -> store.

Embedding happens before anything is written, and a failed Pinecone upsert
rolls the database rows back, so a document either is fully searchable or
does not exist (previously failures left documents with no vectors).
"""
import logging
import os
import re
import uuid
from pathlib import Path
from typing import Any, Dict, List, Tuple

from django.db import transaction

from .embeddings import EmbeddingError, embed_texts
from .llm import detect_domain_llm
from .models import DocumentChunk, UploadedDocument
from .retrieval import delete_chunks, upsert_chunks
from .vectorstore import VectorStoreError

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", "900"))        # characters; ~200 tokens for MiniLM
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", "150"))
OCR_MAX_PAGES = int(os.getenv("OCR_MAX_PAGES", "30"))
MAX_CHUNKS = int(os.getenv("MAX_CHUNKS_PER_DOCUMENT", "1500"))


class IndexingError(Exception):
    """Raised with a user-facing message when a document can't be indexed."""


# ------------------------------ text extraction ------------------------------

def extract_text_from_pdf(path: str) -> List[Tuple[int, str]]:
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            pages = [(i + 1, pg.extract_text() or "") for i, pg in enumerate(pdf.pages)]
        if any(t.strip() for _, t in pages):
            return pages
    except Exception as e:
        logger.warning("pdfplumber failed: %s", e)

    try:
        from pypdf import PdfReader
        pages = [(i + 1, p.extract_text() or "") for i, p in enumerate(PdfReader(path).pages)]
        if any(t.strip() for _, t in pages):
            return pages
    except Exception as e:
        logger.warning("pypdf failed: %s", e)

    # Scanned PDF: OCR (slow, so capped)
    try:
        from pdf2image import convert_from_path
        import pytesseract
        logger.info("Falling back to OCR for %s", path)
        images = convert_from_path(path, dpi=200, last_page=OCR_MAX_PAGES)
        return [(i + 1, pytesseract.image_to_string(img)) for i, img in enumerate(images)]
    except Exception as e:
        logger.warning("OCR failed: %s", e)
    return []


def extract_text_from_docx(path: str) -> List[Tuple[int, str]]:
    import docx  # python-docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [(1, "\n".join(parts))]


def extract_text_from_txt(path: str) -> List[Tuple[int, str]]:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return [(1, f.read())]


def extract_text(path: str) -> List[Tuple[int, str]]:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    if ext == ".docx":
        return extract_text_from_docx(path)
    if ext in (".txt", ".md"):
        return extract_text_from_txt(path)
    raise IndexingError(f"Unsupported file type {ext!r}. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")


# ------------------------------ sectioning/chunking --------------------------

def detect_sections(page_text: str) -> List[Tuple[str, str]]:
    lines = page_text.splitlines()
    headings = []
    for i, line in enumerate(lines):
        s = line.strip()
        if not s or len(s) > 120:
            continue
        if (s.isupper() and len(s) > 3) or re.match(r"^\d+(\.\d+)*[.)]?\s+[A-Z]\S*", s):
            headings.append((i, s))
    if not headings:
        return [("", page_text.strip())]

    sections = []
    if headings[0][0] > 0:
        preamble = "\n".join(lines[:headings[0][0]]).strip()
        if preamble:
            sections.append(("", preamble))
    for idx, (line_idx, title) in enumerate(headings):
        end = headings[idx + 1][0] if idx + 1 < len(headings) else len(lines)
        sections.append((title, "\n".join(lines[line_idx:end]).strip()))
    return sections


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """Split on sentence/word boundaries into ~size-character chunks with overlap."""
    clean = re.sub(r"\s+", " ", text or "").strip()
    if not clean:
        return []
    if len(clean) <= size:
        return [clean]

    chunks, start = [], 0
    while start < len(clean):
        end = min(start + size, len(clean))
        if end < len(clean):
            # Prefer to cut at a sentence end, else at a space, in the last 30% of the window.
            window = clean[start:end]
            cut = max(window.rfind(". "), window.rfind("? "), window.rfind("! "))
            if cut < size * 0.7:
                cut = window.rfind(" ")
            if cut >= size * 0.7:
                end = start + cut + 1
        chunks.append(clean[start:end].strip())
        if end >= len(clean):
            break
        next_start = max(end - overlap, start + 1)
        space = clean.find(" ", next_start)
        start = space + 1 if 0 <= space < end else next_start
    return [c for c in chunks if c]


def normalize_domain(raw: str) -> str:
    r = re.sub(r"[^a-z0-9 _]+", "", (raw or "").lower()).strip().replace(" ", "_")
    return "general" if (not r or r.isdigit() or len(r) < 3) else r[:50]


# ----------------------------------- indexing --------------------------------

def index_document(file_path: str, user_id: int, original_name: str) -> Dict[str, Any]:
    pages = extract_text(file_path)
    if not any(t.strip() for _, t in pages):
        raise IndexingError("No text could be extracted from this file.")

    raw_text = "\n".join(t for _, t in pages)
    domain = normalize_domain(detect_domain_llm(raw_text))
    title = Path(original_name).stem[:500]
    document_id = uuid.uuid4()

    records: List[Dict[str, Any]] = []
    for page_num, page_text in pages:
        for s_idx, (section_title, sec_text) in enumerate(detect_sections(page_text)):
            for c in chunk_text(sec_text):
                records.append({"id": str(uuid.uuid4()), "page_num": page_num, "section_idx": s_idx,
                                "section_title": section_title[:500], "text": c})
    if not records:
        raise IndexingError("No text could be extracted from this file.")
    if len(records) > MAX_CHUNKS:
        raise IndexingError(f"Document is too large ({len(records)} chunks, limit {MAX_CHUNKS}).")

    try:
        vectors = embed_texts([r["text"] for r in records])
    except EmbeddingError as e:
        logger.error("index_document: %s", e)
        raise IndexingError("The embedding model is unavailable; please try again later.") from e

    with transaction.atomic():
        doc = UploadedDocument.objects.create(id=document_id, user_id=user_id, title=title,
                                              original_filename=original_name, domain=domain)
        DocumentChunk.objects.bulk_create([
            DocumentChunk(id=r["id"], document=doc, chunk_idx=i, section_idx=r["section_idx"],
                          section_title=r["section_title"], page_num=r["page_num"],
                          snippet=r["text"][:300], pinecone_id=r["id"])
            for i, r in enumerate(records)
        ])
        items = [{
            "id": r["id"],
            "values": vec,
            "metadata": {
                "user_id": str(user_id),
                "document_id": str(document_id),
                "domain": domain,
                "title": title,
                "page_num": r["page_num"],
                "section_idx": r["section_idx"],
                "section_title": r["section_title"],
                "chunk_idx": i,
                "snippet": r["text"][:300],
                "chunk_text": r["text"],
            },
        } for i, (r, vec) in enumerate(zip(records, vectors))]
        try:
            upsert_chunks(items)
        except VectorStoreError as e:
            logger.error("index_document: %s", e)
            try:
                delete_chunks([r["id"] for r in records])  # remove any batches that did land
            except Exception:
                logger.exception("cleanup after failed upsert also failed")
            raise IndexingError("Could not store the document in the vector index; please try again.") from e

    logger.info("Indexed %s (%d chunks, domain=%s) for user %s", original_name, len(records), domain, user_id)
    return {"document_id": str(document_id), "chunks_indexed": len(records), "domain": domain, "title": title}
